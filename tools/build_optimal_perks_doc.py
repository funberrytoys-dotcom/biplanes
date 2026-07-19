from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DOCX = ROOT / "docs" / "zabeg-perks-optimal.docx"

INK = RGBColor(31, 36, 41)
MUTED = RGBColor(91, 96, 104)
BLUE = RGBColor(34, 73, 112)
BRASS = RGBColor(158, 105, 34)
RED = RGBColor(132, 42, 48)


def set_run(run, size=10.5, bold=False, color=INK, italic=False):
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    run._element.rPr.rFonts.set(qn("w:cs"), "Arial")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def border(cell, color="D8C7AB"):
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        el = borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "6")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)


def para(doc, text="", size=10.5, bold=False, color=INK, after=6, before=0, align=None, italic=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.12
    if align is not None:
        p.alignment = align
    if text:
        set_run(p.add_run(text), size=size, bold=bold, color=color, italic=italic)
    return p


def heading(doc, text, level=1):
    p = para(doc, text, size=17 if level == 1 else 13, bold=True, color=BLUE if level == 1 else BRASS, after=6, before=12)
    if level == 1:
        p_pr = p._p.get_or_add_pPr()
        borders = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "8")
        bottom.set(qn("w:space"), "5")
        bottom.set(qn("w:color"), "D6A24A")
        borders.append(bottom)
        p_pr.append(borders)
    return p


def bullet(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.23)
    p.paragraph_format.first_line_indent = Inches(-0.13)
    p.paragraph_format.space_after = Pt(3)
    set_run(p.add_run("- " + text), size=10, color=INK)


def make_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        shade(cell, "E8D9BD")
        border(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_run(p.add_run(header), size=8.6, bold=True, color=BLUE)
        if widths:
            cell.width = Inches(widths[i])
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            border(cells[i])
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cells[i].paragraphs[0]
            set_run(p.add_run(value), size=8.35, bold=(i == 0), color=INK)
            if widths:
                cells[i].width = Inches(widths[i])
    return table


SOV = [
    ("Раскрутка", "Пулемёт / Деталь", "Первый вкус С.О.В.: темп растёт, если держишь цель.", "Оставить"),
    ("Спаренный ствол", "Пулемёт / Деталь", "Сразу видимый веер пуль, проще попадать на скорости.", "Оставить"),
    ("Трассирующий замок", "Пулемёт / Модуль", "Учит держать одну цель носом, а не просто поливать экран.", "Оставить"),
    ("Радиатор-турбина", "Пулемёт / Модуль", "Связывает скорость самолёта и темп оружия.", "Оставить"),
    ("Гатлинг-шторм", "Пулемёт / Легенда", "Главная корона скорострела.", "Оставить"),
    ("Осиный залп", "Ракеты / Деталь", "Показывает отличие С.О.В.: много микро-ракет.", "Оставить"),
    ("Осиное гнездо", "Звено / Деталь", "Открывает рой дронов без перегруза ИИ.", "Оставить"),
    ("Жало роя", "Звено / Модуль", "Связывает пулемётную метку и дронов.", "Оставить"),
    ("Ремонтный дрон", "Звено / Модуль", "Даёт выживание в стиле С.О.В., не броню.", "Оставить"),
    ("Полный рой", "Звено / Легенда", "Вторая корона фракции, fantasy 'маленькая армия'.", "Оставить"),
    ("Перезалп", "Пилотаж / Модуль", "Ракеты за чистый выход из срыва.", "Оставить"),
    ("Бритва", "Пилотаж / Модуль", "С.О.В. получает лёгкость и точный довираж.", "Оставить"),
]

JACKALS = [
    ("Домна", "Пулемёт / Узел", "Главная механика Шакалов: жар как ресурс.", "Оставить"),
    ("Перегрев = залп", "Пулемёт / Узел", "Переворачивает перегрев из штрафа в риск-награду.", "Оставить"),
    ("Двустволка-картечь", "Пулемёт / Модуль", "Даёт ранний вау-эффект огнешара и луж.", "Оставить"),
    ("Гром-пушка", "Пулемёт / Легенда", "Корона тяжёлого удара.", "Оставить"),
    ("Чугунный гостинец", "Ракеты / Узел", "Одна тяжёлая ракета вместо роя.", "Оставить"),
    ("Напалмовый ковёр", "Ракеты / Модуль", "Маршрут полёта становится огненной линией.", "Оставить"),
    ("Ипритный колокол", "Ракеты / Модуль", "Клеймо зоны связывает ракеты, звено и огонь.", "Оставить"),
    ("Звено асов", "Звено / Узел", "Фракционная альтернатива дронам: мало, живуче, смертно.", "Оставить"),
    ("Огневой приказ", "Звено / Узел", "Решает слабость звена: короткий фокус по цели.", "Оставить"),
    ("Братская броня", "Звено / Модуль", "Эмоция потери аса и защиты от команды.", "Оставить"),
    ("Броненос", "Корпус / Модуль", "Тяжёлый самолёт как стиль, с честной ценой скорости.", "Оставить"),
    ("Таранный киль", "Корпус / Узел", "Самая отличимая шакалья утилита: лобовой риск.", "Оставить"),
]

COMMON = [
    ("Бронепластины", "Корпус / Деталь", "Базовая выживаемость для любого стиля.", "Оставить"),
    ("Полевой ремонт", "Корпус / Модуль", "Лечит только вне давления, не отменяет ошибку.", "Оставить"),
    ("Феникс", "Корпус / Легенда", "Детерминированный второй шанс без рандома.", "Оставить"),
    ("Чистый срыв", "Пилотаж / Деталь", "Сразу завязывает систему на stall.", "Оставить"),
    ("Эгида", "Пилотаж / Модуль", "Защита за чистый выход, а не пассивный щит.", "Оставить"),
    ("Дымзавеса", "Пилотаж / Модуль", "Даёт дыхание и обман lock, но не урон.", "Оставить"),
    ("Мёртвая петля", "Пилотаж / Узел", "Акробатика превращается в перезарядку.", "Оставить"),
    ("Захват целей", "Ракеты / Легенда", "Одна зрелищная общая коронка, но поздняя.", "Оставить осторожно"),
]

DEFER = [
    ("Бесконечная лента / Третий ствол / Перехлёст очередей", "С.О.В.", "Слишком много промежуточных пулемётных ступеней для первого теста."),
    ("Гатлинг-болванка / Гром-болванка / Гром-огнешар", "Сила", "Дублируют роль урон-кирпичей; пока хватит базовых корон."),
    ("Погребальный костёр / Захват целей Шакалов", "Шакалы", "Крутые, но сильно усложняют зоны, клейма и босс-баланс."),
    ("Полное звено Чико", "Шакалы", "Слишком большой ИИ-бласт-радиус для первого внедрения звена."),
    ("Слётанность / Пике-окно / Иммельман-возмездие", "Общие", "Сильные, но требуют точного детектора манёвров."),
    ("Облегчённый каркас", "Общие", "Опасен: может рано отменить цену stall и корпуса."),
]


def build():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)

    style = doc.styles["Normal"]
    style.font.name = "Arial"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    style._element.rPr.rFonts.set(qn("w:cs"), "Arial")
    style.font.size = Pt(10.5)
    style.font.color.rgb = INK

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(footer.add_run("Biplanes / оптимальный первый пул перков"), size=8, color=MUTED)

    para(doc, "BIPLANES", size=24, bold=True, color=BLUE, align=WD_ALIGN_PARAGRAPH.CENTER, after=0)
    para(doc, "Оптимальный первый пул перков «Забега»", size=17, bold=True, color=BRASS, align=WD_ALIGN_PARAGRAPH.CENTER, after=4)
    para(doc, "Решение: 32 перка всего / 20 карт в пуле конкретного забега", size=11, color=MUTED, align=WD_ALIGN_PARAGRAPH.CENTER, after=12)

    heading(doc, "Короткий вердикт")
    para(
        doc,
        "Оптимальное количество для первой полноценной версии системы - 32 перка: 12 С.О.В., 12 Шакалов и 8 общих. "
        "В конкретном забеге игрок видит не все 32, а около 20 карт: 12 своей фракции + 8 общих. "
        "Это даёт достаточно вариантов для 14 выборов, но не превращает первый прототип в балансировочную мясорубку.",
        size=11,
        bold=True,
        color=RED,
    )
    for item in [
        "58 перков оставить как полную дизайн-библию и контент-план, но не как первый объём внедрения.",
        "32 перка сохраняют главное: различие фракций, три стартовых настроения, stall-пилотаж и честные слабости.",
        "Пул на один забег должен быть около 20 карт: меньше - сборки повторяются, больше - игрок перестаёт понимать систему.",
    ]:
        bullet(doc, item)

    heading(doc, "Почему не 58 сразу")
    make_table(
        doc,
        ["Риск", "Что ломается", "Как режем"],
        [
            ("Перегруз выбора", "Игрок не отличит важные карты от декоративных.", "Оставляем только карты, которые меняют стиль полёта."),
            ("Баланс", "14 выборов не способны честно раскрыть 58 карт.", "Делаем пул 20 карт на забег."),
            ("Физика теряется", "Зоны, рои и коронки могут начать играть вместо пилота.", "Сохраняем stall и нос самолёта как главный фильтр."),
            ("Слишком много ИИ", "Звено и дроны быстро становятся отдельной игрой.", "В первом пуле по одной понятной линии звена на фракцию."),
        ],
        widths=[1.7, 2.35, 2.45],
    )

    heading(doc, "Структура первого пула")
    make_table(
        doc,
        ["Блок", "Количество", "Что проверяет"],
        [
            ("С.О.В.", "12", "Темп, точность, микро-ракеты, рой дронов, лёгкий пилотаж."),
            ("Шакалы", "12", "Жар, тяжёлый удар, огненные зоны, асы, броня-таран."),
            ("Общие", "8", "Корпус, ремонт, второй шанс, stall-утилита, одна общая коронка."),
            ("На один забег", "20", "12 фракционных + 8 общих."),
            ("Выборов за забег", "14", "Можно собрать 1 корону + гибрид, но не всё сразу."),
        ],
        widths=[1.55, 1.2, 3.75],
    )

    doc.add_page_break()
    heading(doc, "С.О.В. - оставить 12")
    make_table(doc, ["Перк", "Тип", "Почему в первом пуле", "Решение"], SOV, widths=[1.5, 1.35, 2.95, 0.7])
    doc.add_page_break()

    heading(doc, "Шакалы - оставить 12")
    make_table(doc, ["Перк", "Тип", "Почему в первом пуле", "Решение"], JACKALS, widths=[1.5, 1.35, 2.95, 0.7])

    heading(doc, "Общие - оставить 8")
    make_table(doc, ["Перк", "Тип", "Почему в первом пуле", "Решение"], COMMON, widths=[1.5, 1.35, 2.95, 0.7])
    doc.add_page_break()

    heading(doc, "Что отложить во второй заход")
    para(doc, "Это не выбрасывание идей. Это защита первого плейтеста: сначала проверяем, что система вообще цепляет.", size=10.5)
    make_table(doc, ["Отложить", "Блок", "Причина"], DEFER, widths=[2.35, 1.15, 3.0])

    heading(doc, "Как это ощущается игроком")
    for item in [
        "С.О.В.: быстрый, точный, синий, ройный. Игрок учится держать цель, не отпускать гашетку, играть скоростью.",
        "Шакалы: тяжёлые, горячие, рискованные. Игрок ведёт жар, ставит огонь, тараном покупает момент силы.",
        "Общие: не дают новый урон-класс, а позволяют выжить и превратить stall-манёвр в тактическую кнопку.",
    ]:
        bullet(doc, item)

    heading(doc, "Решение для производства иконок")
    para(
        doc,
        "Для иконок оптимально рисовать сразу эти 32 как основной игровой набор, а оставшиеся 26 держать как backlog. "
        "Так мы не тратим арт-время на карты, которые могут выпасть из первого баланса.",
        size=10.5,
    )
    make_table(
        doc,
        ["Приоритет", "Иконки", "Зачем"],
        [
            ("P0", "32 выбранных", "Нужны для первого рабочего забега и UI."),
            ("P1", "26 отложенных", "Рисовать после плейтеста, когда ясно, какие ветки реально живут."),
            ("Contact sheet", "1 общий лист", "Проверить, что фракции различаются с первого взгляда."),
        ],
        widths=[1.4, 1.6, 3.5],
    )

    heading(doc, "Итог")
    para(
        doc,
        "Мой выбор: 32 перка как оптимальный первый пул. Это достаточно богато, чтобы почувствовать новую игру, "
        "и достаточно узко, чтобы не утопить физику Biplanes под контентом. После плейтеста расширяться до 44, затем до 58.",
        size=11,
        bold=True,
        color=RED,
    )

    DOCX.parent.mkdir(parents=True, exist_ok=True)
    doc.save(DOCX)
    print(DOCX)


if __name__ == "__main__":
    build()
