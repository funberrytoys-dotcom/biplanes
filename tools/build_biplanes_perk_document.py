from __future__ import annotations

import math
import shutil
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont, ImageFilter


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "output" / "biplanes-perks"
TMP_DIR = ROOT / "tmp" / "biplanes-perks"
COVER_SRC = Path.home() / ".codex" / "generated_images" / "019f10a5-f187-7993-89d5-d72c187a7e2f" / "ig_0bc06ce3d54930a9016a41b4b3e4408191b5e59fffa6eb8f1e.png"
COVER_DST = OUT_DIR / "biplanes-perks-cover.png"
DOCX_PATH = OUT_DIR / "Biplanes_Core_Run_Perks.docx"


INK = RGBColor(31, 36, 41)
MUTED = RGBColor(93, 98, 105)
NAVY = RGBColor(28, 58, 85)
BRASS = RGBColor(168, 116, 48)
BURGUNDY = RGBColor(122, 37, 45)
SKY = RGBColor(69, 121, 158)
GREEN = RGBColor(73, 111, 84)
PAPER = "F7F1E6"


CATEGORY = {
    "Штурм": {"color": (156, 50, 45), "accent": "B7322D", "short": "S", "shape": "tracer"},
    "Засев": {"color": (202, 113, 36), "accent": "C97124", "short": "Z", "shape": "bomb"},
    "Звено": {"color": (54, 126, 168), "accent": "367EA8", "short": "V", "shape": "wing"},
    "Корпус": {"color": (83, 128, 84), "accent": "538054", "short": "K", "shape": "shield"},
    "Сила": {"color": (186, 145, 56), "accent": "BA9138", "short": "P", "shape": "engine"},
    "Пилотаж": {"color": (104, 82, 150), "accent": "685296", "short": "M", "shape": "loop"},
}


@dataclass
class Perk:
    name: str
    category: str
    rarity: str
    human: str
    effect: str
    synergy: str
    advice: str
    icon: str


PERKS: list[Perk] = [
    Perk("Пикирующий заряд", "Пилотаж", "Деталь", "Первый залп после пикирования бьёт злее.", "+35% урона и +20% темпа на 1.2 сек после набора скорости в пикировании.", "Гром-пушка, Хвостовой замок", "Отдай высоту, забери темп атаки.", "dive"),
    Perk("Нос вниз", "Пилотаж", "Деталь", "Выход из срыва превращается в рывок.", "После выхода из stall самолёт получает короткий прирост скорости.", "Форсажный подхват", "Не отменяет stall, а награждает за чистое спасение.", "down"),
    Perk("Срывной манёвр", "Пилотаж", "Модуль", "На грани сваливания нос доворачивает резче.", "При скорости около порога stall поворот становится острее.", "Дымная спираль, Огненный след", "Слишком долго держишь нос вверх - сорвёшься.", "stall"),
    Perk("Хвостовой замок", "Штурм", "Модуль", "Чем дольше враг перед носом, тем точнее очередь.", "Удержание цели в секторе атаки копит точность и крит.", "Бронебойные ленты, Охотничий угол", "Перк для настоящего dogfight, не для хаоса.", "target"),
    Perk("Форсаж затвора", "Штурм", "Деталь", "Пулемёт стреляет быстрее.", "+20% темпа огня.", "Гатлинг", "Лучше всего, если часто держишь врага носом.", "bolt"),
    Perk("Бронебойные ленты", "Штурм", "Деталь", "Пули прошивают строй.", "Каждая N-я пуля пробивает следующую цель.", "Перфоратор, Хвостовой замок", "Сильнее против плотных волн.", "ammo"),
    Perk("Спаренный ствол", "Штурм", "Модуль", "Очередь становится шире и увереннее.", "+1 параллельная линия огня с меньшим уроном.", "Гатлинг, Свинцовый шторм", "Прощает небольшую ошибку носа.", "twin"),
    Perk("Гром-пушка", "Штурм", "Узел", "Один заход - один тяжёлый удар.", "Редкий снаряд с высоким уроном и отдачей.", "Пикирующий заряд, Тяжёлый спуск", "Промах больнее, чем у пулемёта.", "cannon"),
    Perk("Гатлинг", "Штурм", "Легенда", "Нос самолёта превращается в пилу.", "Очень высокий темп, растущий при удержании цели.", "Форсаж затвора + Спаренный ствол", "Сильная легенда, но требует позиции.", "gatling"),
    Perk("Бомболюк", "Засев", "Деталь", "Самолёт начинает оставлять смерть под собой.", "Открывает бомбы или усиливает частоту сброса.", "Напалмовый ковёр", "Высота становится ресурсом.", "bombbay"),
    Perk("Серийный сброс", "Засев", "Деталь", "Одна команда - несколько бомб по дуге.", "+2 малые бомбы в серии.", "Каскадный подрыв", "Работает, если заранее строишь маршрут.", "cluster"),
    Perk("Зажигательная смесь", "Засев", "Модуль", "После взрыва небо ещё горит.", "Бомбы и ракеты оставляют короткую огненную зону.", "Напалмовый ковёр", "Лучше против роёв, хуже против одиночного аса.", "fire"),
    Perk("Минный сев", "Засев", "Модуль", "Позади Чико остаются воздушные капканы.", "Мины висят несколько секунд и взрываются при сближении.", "Небесная паутина", "Слабый старт, сильное поле к середине боя.", "mine"),
    Perk("Ипритный колокол", "Засев", "Узел", "Облако делает врагов мягче для всего оружия.", "Зона помечает цели: они получают больше урона.", "Бомбы + дроны", "Главный клей гибридного контроля.", "gas"),
    Perk("Напалмовый ковёр", "Засев", "Легенда", "Бомбы рисуют длинную огненную трассу.", "Горящая зона тянется вдоль траектории сброса.", "Бомболюк + Зажигательная смесь", "Заставляет думать линией полёта.", "napalm"),
    Perk("Ведомый Птенец", "Звено", "Деталь", "У хвоста появляется маленький помощник.", "Один ведомый стреляет вместе с Чико.", "Эскадрилья", "Не должен играть вместо игрока.", "wingman"),
    Perk("Спарка ведомого", "Звено", "Деталь", "Ведомый получает второй ствол.", "+урон ведомого и лучшая синхронизация залпа.", "Эскадрилья", "Сильнее при плавном пилотаже.", "winggun"),
    Perk("Ведущий поток", "Звено", "Модуль", "Звено держится лучше, если Чико летит чисто.", "Плавный полёт без резких потерь скорости усиливает ведомых.", "Хвостовой ас", "Хаотичный полёт ломает строй.", "flow"),
    Perk("Карусель", "Звено", "Узел", "Дроны кружат вокруг самолёта и режут ближнюю зону.", "Часть ведомых переходит в защитную орбиту.", "Корпус, мины", "Безопасно, но хуже фокусирует босса.", "orbit"),
    Perk("Эскадрилья", "Звено", "Легенда", "Ведомые собираются в боевой клин.", "Несколько ведомых стреляют залпом вместе с Чико.", "Птенец + Спарка ведомого", "Про летать самому, но не одному.", "squad"),
    Perk("Усиление корпуса", "Корпус", "Деталь", "Самолёт дольше держит бой.", "+25% прочности.", "Эгида", "Клей билда, не победная кнопка.", "plate"),
    Perk("Планёр", "Корпус", "Деталь", "Сваливание становится чуть мягче.", "Падение при stall медленнее, но stall остаётся опасным.", "Срывник, Звено", "Прощает ошибку, но не отменяет физику.", "glider"),
    Perk("Низовой рывок", "Корпус", "Модуль", "У нижней зоны форсаж вытаскивает сильнее.", "Форсаж эффективнее близко к нижней границе арены.", "Ремонт, Нырок", "Если пришёл туда в штопоре - поздно.", "lowboost"),
    Perk("Эгида", "Корпус", "Узел", "Бронеплита обновляется после чистого манёвра.", "Периодический щит после безопасного выхода из опасной зоны.", "Планёр + Усиление корпуса", "Защита за пилотаж, не за стояние.", "aegis"),
    Perk("Феникс", "Корпус", "Легенда", "Один раз самолёт держится там, где должен был погибнуть.", "Один раз за забег смертельный удар оставляет 1 HP и гасит двигатель.", "Корпус + Пилотаж", "Драматичное спасение без случайности.", "phoenix"),
    Perk("Овердрайв", "Сила", "Узел", "Весь арсенал работает быстрее и злее.", "+урон и +темп для пуль, бомб и ведомых.", "Гибридные билды", "Лучший клей, если собрал несколько источников урона.", "overdrive"),
    Perk("Красная зона", "Сила", "Модуль", "Мотор орёт, но держит ещё одну атаку.", "При перегреве остаётся короткое окно усиленного форсажа.", "Сердце мотора", "После окна двигатель дольше мёртвый.", "redline"),
    Perk("Мёртвая петля", "Пилотаж", "Узел", "Опасная петля сама перезаряжает спецоружие.", "Полный разворот через верх ускоряет перезарядку ракет/бомб.", "Ракеты, Гром-пушка", "В верхней точке легко потерять скорость.", "loop"),
    Perk("Дымная подпись", "Засев", "Модуль", "Копоть рисует ложный путь.", "После резкого разворота остаётся дымовая завеса, сбивающая прицел врагов.", "Срывной манёвр", "Не урон, а пространство для дыхания.", "smoke"),
    Perk("Пока горят огни", "Сила", "Легенда", "Маршрут ещё не спасён, значит самолёт держится.", "На финальной волне усиливает выбранный главный стиль билда.", "Любая собранная легенда", "Финальный героический капстоун, не ранний snowball.", "beacon"),
]


BUILDS = [
    ("Штурмовик", "Носовая мясорубка", "Держать врага перед носом и превращать заход в непрерывную очередь.", ["Гатлинг", "Гром-пушка", "Огненный штурм"], "Проседает, если игрок не умеет держать курс на цель."),
    ("Дуэлянт", "Один заход - один труп", "Играть через хвост, пикирование, точный угол и короткие окна атаки.", ["Бронебойный ас", "Пикирующий казначей", "Охотник на босса"], "Слабее против роя без добора зоны или звена."),
    ("Бомбардир", "Выжженное небо", "Побеждать маршрутом: бомбы, огонь, облака, каскадные зоны.", ["Напалмовый ковёр", "Ипритный колокол", "Каскадный подрыв"], "Нужна траектория, хаотичный полёт кладёт зоны впустую."),
    ("Минёр", "Небесная паутина", "Заранее засевать воздух ловушками и заставлять врагов влетать в поле.", ["Воздушные мины", "Рой ос", "Капкан для босса"], "Медленный старт, пока поле не построено."),
    ("Командир звена", "Маленькая эскадрилья", "Летать самому, но растить рядом помощников, которые усиливают стиль.", ["Эскадрилья", "Карусель", "Тётушка-флагман"], "Размазанный урон, нужен узел для фокуса босса."),
    ("Бронекрейсер", "Неубиваемый биплан", "Держать бой, выходить из ошибок и доклеивать активный источник урона.", ["Эгида", "Овердрайв", "Феникс"], "Сам по себе не выигрывает, ему нужен урон."),
]


SYNERGIES = [
    ("Пикирующий заряд + Гром-пушка", "Тяжёлый заход сверху", "Высота становится боеприпасом: отдал её, получил мощный снаряд."),
    ("Срывной манёвр + Дымная подпись", "Опасная петля", "Игрок режет медленные развороты, оставляя ложные пути и зоны."),
    ("Хвостовой замок + Бронебойные ленты", "Dogfight-пробой", "Позиционная игра награждает не только уроном, но и прошивом роя."),
    ("Форсажный подхват + Ракеты", "Спасся и наказал", "Выход из stall сразу превращается в контратаку."),
    ("Высотный запас + Бомболюк", "Вертикальный контроль", "Чем умнее держишь высоту, тем лучше управляешь ареной."),
    ("Ведущий поток + Эскадрилья", "Чистый строй", "Ведомые сильны не числом, а плавным пилотажем игрока."),
    ("Минный сев + Ипритный колокол", "Воздушная ловушка", "Зона помечает врагов, мины добирают тех, кто полез внутрь."),
    ("Планёр + Срывник", "Прощённая ошибка", "Можно дольше играть рядом со stall, но нельзя отменить его цену."),
    ("Овердрайв + любой гибрид", "Клей арсенала", "Пули, бомбы и звено начинают масштабироваться вместе."),
    ("Пока горят огни + легенда ветки", "Финальный почерк", "Последняя волна усиливает именно то, кем стал самолёт в этом забеге."),
]


def font(size: int, bold: bool = False) -> ImageFont.FreeTypeFont:
    candidates = [
        Path("C:/Windows/Fonts/segoeuib.ttf" if bold else "C:/Windows/Fonts/segoeui.ttf"),
        Path("C:/Windows/Fonts/arialbd.ttf" if bold else "C:/Windows/Fonts/arial.ttf"),
        Path("C:/Windows/Fonts/calibrib.ttf" if bold else "C:/Windows/Fonts/calibri.ttf"),
    ]
    for candidate in candidates:
        if candidate.exists():
            return ImageFont.truetype(str(candidate), size)
    return ImageFont.load_default()


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_border(cell, color="D8C7AB", size="8") -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def paragraph_border_bottom(paragraph, color="A87430", size="12") -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        p_pr.append(borders)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), "6")
    bottom.set(qn("w:color"), color)
    borders.append(bottom)


def set_run(run, size=None, bold=None, color=None, italic=None) -> None:
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    run._element.rPr.rFonts.set(qn("w:cs"), "Arial")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def add_para(doc, text="", size=10.5, bold=False, color=INK, after=6, before=0, italic=False, align=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.12
    if align is not None:
        p.alignment = align
    if text:
        r = p.add_run(text)
        set_run(r, size=size, bold=bold, color=color, italic=italic)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14 if level == 1 else 9)
    p.paragraph_format.space_after = Pt(6)
    r = p.add_run(text)
    set_run(r, size=17 if level == 1 else 13, bold=True, color=NAVY if level == 1 else BRASS)
    if level == 1:
        paragraph_border_bottom(p, "D6A24A", "8")
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style=None)
    p.paragraph_format.left_indent = Inches(0.22)
    p.paragraph_format.first_line_indent = Inches(-0.12)
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run("- " + text)
    set_run(r, size=10, color=INK)
    return p


def draw_icon(perk: Perk, path: Path) -> None:
    info = CATEGORY[perk.category]
    base = Image.new("RGB", (520, 520), "#" + PAPER)
    draw = ImageDraw.Draw(base)
    color = info["color"]
    dark = tuple(max(0, c - 55) for c in color)
    light = tuple(min(255, c + 65) for c in color)

    for y in range(520):
        t = y / 520
        r = int((247 * (1 - t)) + (225 * t))
        g = int((241 * (1 - t)) + (226 * t))
        b = int((230 * (1 - t)) + (205 * t))
        draw.line((0, y, 520, y), fill=(r, g, b))

    draw.rounded_rectangle((20, 20, 500, 500), radius=34, fill=(32, 38, 43), outline=dark, width=8)
    draw.rounded_rectangle((34, 34, 486, 486), radius=26, fill=(48, 54, 58), outline=light, width=2)
    draw.rectangle((34, 360, 486, 486), fill=dark)

    # Sunburst / instrument geometry.
    cx, cy = 260, 225
    for angle in range(0, 360, 20):
        rad = math.radians(angle)
        x = cx + math.cos(rad) * 190
        y = cy + math.sin(rad) * 190
        draw.line((cx, cy, x, y), fill=(74, 78, 80), width=2)
    draw.ellipse((90, 55, 430, 395), fill=(61, 67, 70), outline=(220, 185, 108), width=5)
    draw.ellipse((125, 90, 395, 360), outline=(102, 106, 104), width=4)

    # Category emblem.
    shape = info["shape"]
    if shape == "tracer":
        for off in (-46, 0, 46):
            draw.line((130, 260 + off, 390, 155 + off), fill=light, width=14)
            draw.polygon([(390, 155 + off), (350, 145 + off), (365, 185 + off)], fill=(240, 208, 92))
    elif shape == "bomb":
        for x in (200, 260, 320):
            draw.ellipse((x - 24, 140, x + 24, 232), fill=(68, 73, 75), outline=light, width=4)
            draw.polygon([(x - 20, 145), (x + 20, 145), (x, 108)], fill=dark)
            draw.line((x, 232, x, 292), fill=(236, 176, 72), width=8)
        draw.arc((120, 245, 400, 380), 195, 345, fill=(240, 112, 45), width=18)
    elif shape == "wing":
        draw.polygon([(125, 245), (260, 130), (395, 245), (315, 265), (260, 210), (205, 265)], fill=light, outline=dark)
        draw.polygon([(160, 305), (260, 235), (360, 305), (295, 318), (260, 285), (225, 318)], fill=(190, 151, 72), outline=dark)
        for x in (170, 260, 350):
            draw.ellipse((x - 16, 285, x + 16, 317), fill=(235, 235, 210))
    elif shape == "shield":
        draw.polygon([(260, 90), (390, 145), (360, 310), (260, 375), (160, 310), (130, 145)], fill=light, outline=dark)
        for x in range(170, 351, 45):
            draw.ellipse((x, 180, x + 16, 196), fill=(40, 45, 45))
        draw.line((180, 270, 340, 180), fill=dark, width=13)
    elif shape == "engine":
        draw.ellipse((130, 95, 390, 355), fill=light, outline=dark, width=8)
        for angle in range(0, 360, 30):
            rad = math.radians(angle)
            x1 = cx + math.cos(rad) * 52
            y1 = cy + math.sin(rad) * 52
            x2 = cx + math.cos(rad) * 120
            y2 = cy + math.sin(rad) * 120
            draw.line((x1, y1, x2, y2), fill=dark, width=10)
        draw.ellipse((210, 175, 310, 275), fill=(245, 232, 154), outline=dark, width=5)
        draw.line((300, 155, 350, 95), fill=(248, 232, 92), width=10)
        draw.line((340, 96, 316, 104), fill=(248, 232, 92), width=10)
    else:
        points = [(130, 260), (170, 120), (315, 120), (390, 210), (330, 315), (210, 340), (150, 280)]
        draw.line(points, fill=light, width=15, joint="curve")
        draw.polygon([(388, 210), (355, 188), (362, 230)], fill=(242, 220, 100))

    # Small custom glyph overlay.
    glyph_font = font(68, True)
    draw.rounded_rectangle((58, 58, 134, 134), radius=16, fill=dark, outline=(235, 209, 139), width=3)
    draw.text((96, 94), info["short"], font=glyph_font, fill=(248, 235, 176), anchor="mm")

    title_font = font(34, True)
    rarity_font = font(24, False)
    wrapped = perk.name
    if len(wrapped) > 17:
        words = wrapped.split()
        if len(words) > 1:
            wrapped = " ".join(words[: len(words) // 2]) + "\n" + " ".join(words[len(words) // 2 :])
    bbox = draw.multiline_textbbox((260, 397), wrapped, font=title_font, anchor="ma", spacing=2)
    draw.multiline_text((260, 397), wrapped, font=title_font, fill=(255, 246, 220), anchor="ma", align="center", spacing=2)
    draw.text((260, 462), f"{perk.category} · {perk.rarity}", font=rarity_font, fill=(228, 207, 151), anchor="mm")

    base = base.filter(ImageFilter.UnsharpMask(radius=1, percent=110, threshold=3))
    path.parent.mkdir(parents=True, exist_ok=True)
    base.save(path)


def make_all_icons() -> dict[str, Path]:
    paths = {}
    icon_dir = OUT_DIR / "perk-icons"
    icon_dir.mkdir(parents=True, exist_ok=True)
    for idx, perk in enumerate(PERKS, 1):
        path = icon_dir / f"{idx:02d}_{slug(perk.name)}.png"
        draw_icon(perk, path)
        paths[perk.name] = path
    return paths


def slug(text: str) -> str:
    table = str.maketrans(
        "абвгдеёжзийклмнопрстуфхцчшщъыьэюя ",
        "abvgdeejzijklmnoprstufhzcss_y_eua_",
    )
    return text.lower().translate(table).replace("__", "_").strip("_")


def setup_doc() -> Document:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.68)
    section.left_margin = Inches(0.72)
    section.right_margin = Inches(0.72)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)
    styles = doc.styles
    for style_name in ("Normal", "Body Text"):
        style = styles[style_name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style._element.rPr.rFonts.set(qn("w:cs"), "Arial")
        style.font.size = Pt(10.5)
        style.font.color.rgb = INK
    return doc


def add_footer(section, text: str) -> None:
    p = section.footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    set_run(r, size=8, color=MUTED)


def add_cover(doc: Document) -> None:
    if COVER_SRC.exists():
        OUT_DIR.mkdir(parents=True, exist_ok=True)
        shutil.copy2(COVER_SRC, COVER_DST)
    p = add_para(doc, "BIPLANES", size=28, bold=True, color=NAVY, after=0, align=WD_ALIGN_PARAGRAPH.CENTER)
    p.paragraph_format.space_before = Pt(4)
    add_para(doc, "Ядро забега: перки, билды, эволюции и синергии", size=17, bold=True, color=BRASS, after=4, align=WD_ALIGN_PARAGRAPH.CENTER)
    add_para(doc, "Ночная редколлегия агентов · версия для обсуждения", size=10, color=MUTED, after=10, align=WD_ALIGN_PARAGRAPH.CENTER)
    if COVER_DST.exists():
        doc.add_picture(str(COVER_DST), width=Inches(7.0))
    add_para(doc, "Короткий вывод", size=13, bold=True, color=NAVY, after=4)
    add_para(
        doc,
        "Лучшее ядро забега для Biplanes - не большой магазин чисел, а компактная система, где каждый перк меняет вопрос игрока: "
        "как мне теперь лететь? Побеждает не абстрактный DPS, а манера полёта: пикировать, держать хвост, рисовать огненные линии, вести звено, играть на грани stall.",
        size=11,
        after=7,
    )
    for item in [
        "Старт забега: мягкий выбор архетипа, который повышает шанс нужных карт, но не запирает свободу.",
        "Внутри забега: 14 выборов, каждый раз 1 из 3 карт; ключевые пороги на 3 и 6 вложениях в стиль.",
        "Главная проверка качества: игрок после забега должен сказать не '+20% урона', а 'я собрал пикирующего дуэлянта' или 'я стал минным дирижёром'.",
    ]:
        add_bullet(doc, item)
    doc.add_page_break()


def add_loop_engineering(doc: Document) -> None:
    add_heading(doc, "Как работала ночная петля", 1)
    add_para(doc, "Я применил подход из гайда по loop engineering: цель, роли, независимые проверки и финальный синтез.", size=10.5)
    table = doc.add_table(rows=1, cols=4)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    headers = ["Роль", "Что проверяла", "Главный вклад", "Голос"]
    for i, h in enumerate(headers):
        c = table.rows[0].cells[i]
        set_cell_shading(c, "E8D9BD")
        set_cell_border(c)
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_run(p.add_run(h), size=9, bold=True, color=NAVY)
    rows = [
        ("Бенчмаркер", "Vampire Survivors, Brotato, Hades, Risk of Rain", "Компактный пул, сильные ветки, спасатели от плохого рандома", "За"),
        ("Физик", "Stall, энергия, высота, форсаж", "Перки должны менять манеру полёта", "За"),
        ("Buildcraft", "Архетипы, слабости, эволюции", "6 билдов и мягкая специализация", "За"),
        ("UX", "Карточки и читаемость", "Деталь / Модуль / Узел / Легенда", "За"),
        ("Арт/лор", "Чико, люмены, pulp-авиация", "Перки как предметы, манёвры и суеверия ангара", "За"),
    ]
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_border(cells[i])
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cells[i].paragraphs[0]
            set_run(p.add_run(value), size=8.5, bold=(i == 0), color=INK if i != 3 else GREEN)
    add_para(doc, "Решение совета", size=12, bold=True, color=BURGUNDY, after=3, before=8)
    add_para(
        doc,
        "Берём модель: мягкие архетипы на старте + компактный пул + синергии через пилотаж. "
        "Не берём полную случайность для первого прототипа и не делаем жёсткие классы самолётов до тех пор, пока ядро боя не доказало кайф.",
        size=10.5,
    )


def add_builds(doc: Document) -> None:
    add_heading(doc, "Шесть стартовых фантазий забега", 1)
    add_para(doc, "Игрок выбирает не класс навсегда, а настроение вылета. Выбор мягко подкручивает выдачу карт и даёт раннюю личность забега.", size=10.5)
    for idx in range(0, len(BUILDS), 2):
        table = doc.add_table(rows=1, cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for j, build in enumerate(BUILDS[idx : idx + 2]):
            name, fantasy, core, variants, weakness = build
            cell = table.rows[0].cells[j]
            set_cell_border(cell, "D2C0A0", "10")
            set_cell_shading(cell, "FAF6ED")
            p = cell.paragraphs[0]
            set_run(p.add_run(name), size=13, bold=True, color=NAVY)
            p.add_run("\n")
            set_run(p.add_run(fantasy), size=10.5, bold=True, color=BRASS)
            p.add_run("\n")
            set_run(p.add_run(core), size=9.3, color=INK)
            p.add_run("\n")
            set_run(p.add_run("Варианты: " + ", ".join(variants)), size=9.0, color=GREEN, bold=True)
            p.add_run("\n")
            set_run(p.add_run("Слабость: " + weakness), size=9.0, color=BURGUNDY)
    add_para(doc, "Правило анти-меты", size=12, bold=True, color=NAVY, before=8, after=3)
    for item in [
        "У каждого билда есть честная дырка: точность, толпа, босс, слабый старт, размазанный урон или нехватка DPS.",
        "Волны должны проверять разные навыки: рой, быстрые асы, бронированные цели, преследование, финальный босс.",
        "За 14 выборов нельзя собрать всё. Игрок строит личность, а не заполняет весь магазин.",
    ]:
        add_bullet(doc, item)


def add_progression_rules(doc: Document) -> None:
    add_heading(doc, "Правила выдачи и эволюций", 1)
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(["Слой", "Как работает", "Зачем игроку"]):
        c = table.rows[0].cells[i]
        set_cell_shading(c, "E8EEF5")
        set_cell_border(c)
        p = c.paragraphs[0]
        set_run(p.add_run(h), size=9, bold=True, color=NAVY)
    rows = [
        ("Стартовый контракт", "Игрок выбирает настроение: Штурм, Засев, Звено, Корпус, Пилотаж.", "Забег сразу получает личность, но не становится рельсой."),
        ("Выбор 1 из 3", "Каждые 20-30 секунд показываем 3 карты.", "Решение короткое и понятное на мобайле."),
        ("Пипки ветки", "Каждая карта даёт очко категории; 3 очка открывают Узел, 6 - Легенду.", "Игрок видит прогресс к сборке."),
        ("Гарантия рецепта", "Если эволюция стала доступна впервые, она появляется в одной из трёх карт.", "Сборка ощущается заслуженной, а не украденной рандомом."),
        ("Контроль рандома", "1 reroll и 1 lock за забег в прототипе.", "Спасает плохую выдачу, но не даёт каждый раз собирать идеал."),
        ("Микс-карта", "В каждой тройке минимум один вариант не из главной ветки.", "Оставляет пространство для неожиданных гибридов."),
    ]
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_border(cells[i])
            p = cells[i].paragraphs[0]
            set_run(p.add_run(value), size=8.6, bold=(i == 0), color=INK)


def add_card_language(doc: Document) -> None:
    add_heading(doc, "Язык карточки", 1)
    add_para(doc, "Карточка должна читаться как маленький авиаплакат: стиль боя, эффект, маршрут к эволюции.", size=10.5)
    table = doc.add_table(rows=1, cols=5)
    for i, h in enumerate(["Категория", "Цвет", "Форма", "Игрок считывает", "Роль"]):
        c = table.rows[0].cells[i]
        set_cell_shading(c, "E8D9BD")
        set_cell_border(c)
        set_run(c.paragraphs[0].add_run(h), size=8.5, bold=True, color=NAVY)
    rows = [
        ("Штурм", "красный/сталь", "трассеры, гильзы", "Бью туда, куда смотрит нос", "главный урон"),
        ("Засев", "оранжевый/зелёный", "бомбы, облака, фитили", "Оставляю опасное небо позади", "контроль арены"),
        ("Звено", "синий/латунь", "ведомые, флажки", "Команда воюет рядом", "помощники"),
        ("Корпус", "зелёный/бронза", "заклёпки, щиты", "Живу дольше и выхожу из ошибок", "выживание"),
        ("Сила", "золото/белый", "мотор, молнии", "Весь арсенал сильнее", "гибридный клей"),
        ("Пилотаж", "фиолетовый/золото", "петли, нос, скорость", "Летаю иначе", "уникальность Biplanes"),
    ]
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_border(cells[i])
            set_run(cells[i].paragraphs[0].add_run(value), size=8.2, bold=(i == 0), color=INK)
    add_para(doc, "Шаблон текста", size=12, bold=True, color=NAVY, before=8, after=3)
    for item in [
        "Название: Форсаж затвора",
        "Метка: Штурм · Деталь · II",
        "Человеческая строка: Пулемёт стреляет быстрее.",
        "Числовой эффект: +20% темпа огня.",
        "Синергия: Ближе к Гатлингу.",
        "Подсказка: Хорошо, если часто держишь врага носом.",
    ]:
        add_bullet(doc, item)


def add_perk_cards(doc: Document, icons: dict[str, Path]) -> None:
    add_heading(doc, "Карточки перков", 1)
    add_para(doc, "Ниже - рабочий пул из 30 карт. Для прототипа можно взять первые 18-24, остальные оставить как расширение.", size=10.5)
    for idx in range(0, len(PERKS), 2):
        if idx > 0 and idx % 4 == 0:
            doc.add_page_break()
        table = doc.add_table(rows=1, cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        for j, perk in enumerate(PERKS[idx : idx + 2]):
            cell = table.rows[0].cells[j]
            set_cell_border(cell, CATEGORY[perk.category]["accent"], "10")
            set_cell_shading(cell, "FBF7EF")
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run()
            run.add_picture(str(icons[perk.name]), width=Inches(1.23))
            p = cell.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_run(p.add_run(perk.name), size=11.4, bold=True, color=NAVY)
            p = cell.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_run(p.add_run(f"{perk.category} · {perk.rarity}"), size=8.4, bold=True, color=BRASS)
            for label, text, color in [
                ("Что чувствует игрок: ", perk.human, INK),
                ("Эффект: ", perk.effect, INK),
                ("Синергия: ", perk.synergy, GREEN),
                ("Риск/подсказка: ", perk.advice, BURGUNDY),
            ]:
                p = cell.add_paragraph()
                p.paragraph_format.space_after = Pt(2)
                set_run(p.add_run(label), size=8.1, bold=True, color=color)
                set_run(p.add_run(text), size=8.1, color=INK)


def add_synergies(doc: Document) -> None:
    add_heading(doc, "Миксы и синергии", 1)
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, h in enumerate(["Комбо", "Название ощущения", "Почему это даёт новый опыт"]):
        c = table.rows[0].cells[i]
        set_cell_shading(c, "E8EEF5")
        set_cell_border(c)
        set_run(c.paragraphs[0].add_run(h), size=8.5, bold=True, color=NAVY)
    for combo, feel, why in SYNERGIES:
        cells = table.add_row().cells
        for i, value in enumerate([combo, feel, why]):
            set_cell_border(cells[i])
            set_run(cells[i].paragraphs[0].add_run(value), size=8.2, bold=(i == 1), color=INK if i != 1 else BRASS)
    add_para(doc, "Что нельзя делать", size=12, bold=True, color=BURGUNDY, before=8, after=3)
    for item in [
        "Не давать автостабилизацию, которая сама спасает от stall.",
        "Не давать автоповорот к цели: самолёт должен стрелять туда, куда игрок сумел направить нос.",
        "Не делать оружие, которое чистит экран независимо от направления полёта.",
        "Не превращать ведомых в бессмертный рой, который играет вместо игрока.",
        "Не делать случайные спасения от смерти: та же команда должна давать тот же результат.",
    ]:
        add_bullet(doc, item)


def add_prototype_scope(doc: Document) -> None:
    add_heading(doc, "Рекомендация для первого прототипа", 1)
    add_para(doc, "Чтобы быстро проверить, «вкатывает ли», не надо сразу делать весь пул. Нужен вертикальный срез, где видны разные манеры полёта.", size=10.5)
    rows = [
        ("Минимум", "18 карт", "3 Штурм, 3 Засев, 3 Звено, 3 Корпус, 4 Пилотаж, 2 Сила"),
        ("Эволюции", "6 штук", "Гатлинг, Гром-пушка, Напалм, Эскадрилья, Эгида, Овердрайв"),
        ("Выдача", "14 выборов", "Порог 3 = Узел, порог 6 = Легенда"),
        ("Контроль", "reroll + lock", "По одному на забег, без покупки силы"),
        ("Проверка", "после забега", "Показать: стиль, урон, зоны, звено, спасения, любимая синергия"),
    ]
    table = doc.add_table(rows=1, cols=3)
    for i, h in enumerate(["Блок", "Объём", "Состав"]):
        c = table.rows[0].cells[i]
        set_cell_shading(c, "E8D9BD")
        set_cell_border(c)
        set_run(c.paragraphs[0].add_run(h), size=8.5, bold=True, color=NAVY)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            set_cell_border(cells[i])
            set_run(cells[i].paragraphs[0].add_run(value), size=8.5, bold=(i == 0), color=INK)
    add_para(doc, "Финальная формула", size=12, bold=True, color=NAVY, before=9, after=3)
    add_para(
        doc,
        "Biplanes должен быть roguelike не про «выбери самый большой процент», а про «сегодня я научился летать иначе». "
        "Каждая карта либо меняет траекторию, либо усиливает понятную траекторию, либо помогает выжить после честной ошибки.",
        size=11,
        bold=True,
        color=BURGUNDY,
    )


def build_doc() -> None:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    TMP_DIR.mkdir(parents=True, exist_ok=True)
    icons = make_all_icons()
    doc = setup_doc()
    add_footer(doc.sections[0], "Biplanes · ядро забега · перки и синергии")
    add_cover(doc)
    add_loop_engineering(doc)
    add_builds(doc)
    doc.add_page_break()
    add_progression_rules(doc)
    add_card_language(doc)
    doc.add_page_break()
    add_perk_cards(doc, icons)
    doc.add_page_break()
    add_synergies(doc)
    doc.add_page_break()
    add_prototype_scope(doc)
    doc.save(DOCX_PATH)


if __name__ == "__main__":
    build_doc()
    print(DOCX_PATH)
